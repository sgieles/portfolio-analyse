"""Excel export — multi-sheet workbook with styled tables."""

from __future__ import annotations

import math
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side,
)
from openpyxl.utils import get_column_letter

from analytics.scenario import run_scenario_analysis
from models.results import AnalysisResult


# ── Style constants ────────────────────────────────────────────────────────────
_BLUE_DARK  = "0F3460"
_BLUE_MID   = "1A3A6B"
_ACCENT     = "E94560"
_ALT_ROW    = "F2F5FA"
_BORDER_COL = "CCCCCC"

_HDR_FONT  = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
_BODY_FONT = Font(name="Calibri", size=9)
_TITLE_FONT = Font(name="Calibri", bold=True, size=14, color=_BLUE_DARK)
_LABEL_FONT = Font(name="Calibri", bold=True, size=10, color=_BLUE_DARK)

_HDR_FILL  = PatternFill("solid", fgColor=_BLUE_DARK)
_ALT_FILL  = PatternFill("solid", fgColor=_ALT_ROW)
_GOOD_FILL = PatternFill("solid", fgColor="E8F5E9")
_BAD_FILL  = PatternFill("solid", fgColor="FFEBEE")

_THIN = Side(border_style="thin", color=_BORDER_COL)
_CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_CENTER = Alignment(horizontal="center", vertical="center")
_RIGHT  = Alignment(horizontal="right",  vertical="center")
_LEFT   = Alignment(horizontal="left",   vertical="center")


def _pct(v: float, dec: int = 2) -> str:
    return "" if math.isnan(v) else f"{v * 100:.{dec}f} %"


def _num(v: float, dec: int = 2) -> str:
    return "" if math.isnan(v) else f"{v:.{dec}f}"


def _eur(v: float) -> str:
    return "" if math.isnan(v) else f"€ {v:,.2f}"


def _hdr_cell(ws, row: int, col: int, value: str) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.font   = _HDR_FONT
    c.fill   = _HDR_FILL
    c.border = _CELL_BORDER
    c.alignment = _CENTER


def _data_cell(
    ws, row: int, col: int, value: str,
    align: Alignment = _LEFT,
    alt: bool = False,
) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.font      = _BODY_FONT
    c.border    = _CELL_BORDER
    c.alignment = align
    if alt:
        c.fill = _ALT_FILL


def _auto_width(ws) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


# ── Sheet builders ─────────────────────────────────────────────────────────────

def _build_summary(ws, result: AnalysisResult) -> None:
    portfolio = result.portfolio
    name = portfolio.name if portfolio else "Portfolio"
    period = portfolio.period if portfolio else "?"
    benchmark = portfolio.benchmark if portfolio else "?"

    ws.title = "Summary"
    ws.sheet_view.showGridLines = False

    # Title block
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = f"Portfolio Analysis Report — {name}"
    title_cell.font  = _TITLE_FONT
    title_cell.alignment = _LEFT

    info_rows = [
        ("Generated", date.today().isoformat()),
        ("Period", period),
        ("Benchmark", benchmark),
    ]
    for i, (label, value) in enumerate(info_rows, start=2):
        ws.cell(row=i, column=1, value=label).font = _LABEL_FONT
        ws.cell(row=i, column=2, value=value).font = _BODY_FONT

    # Metrics table
    row = 6
    ws.merge_cells(f"A{row}:B{row}")
    title = ws.cell(row=row, column=1, value="PORTFOLIO METRICS")
    title.font = _LABEL_FONT
    row += 1

    sections = [
        ("Return", [
            ("Expected Return (ann.)", _pct(result.portfolio_return)),
            ("CAGR",                   _pct(result.portfolio_cagr)),
            ("CAPM Expected Return",   _pct(result.portfolio_expected_return_capm)),
        ]),
        ("Risk", [
            ("Volatility (ann.)",      _pct(result.portfolio_volatility)),
            ("Sharpe Ratio",           _num(result.sharpe_ratio)),
            ("Sortino Ratio",          _num(result.sortino_ratio)),
            ("Beta",                   _num(result.beta)),
            ("Max Drawdown",           _pct(result.max_drawdown)),
            ("VaR 95 %",               _pct(result.var_95)),
            ("VaR 99 %",               _pct(result.var_99)),
            ("CVaR (95 %)",            _pct(result.cvar)),
        ]),
        ("Diversification", [
            ("# Assets",               str(len(result.asset_metrics))),
            ("Avg Correlation",        _num(result.avg_correlation)),
            ("Diversification Score",  _num(result.diversification_score)),
        ]),
    ]
    for sec_name, rows in sections:
        _hdr_cell(ws, row, 1, sec_name)
        _hdr_cell(ws, row, 2, "Value")
        row += 1
        for i, (label, val) in enumerate(rows):
            alt = i % 2 == 1
            _data_cell(ws, row, 1, label, alt=alt)
            _data_cell(ws, row, 2, val, align=_RIGHT, alt=alt)
            row += 1
        row += 1  # blank row between sections

    _auto_width(ws)


def _build_assets(ws, result: AnalysisResult) -> None:
    ws.title = "Assets"
    ws.sheet_view.showGridLines = False

    headers = [
        "Ticker", "Weight", "Latest Price (€)",
        "CAGR", "Return (ann.)", "Volatility",
        "Sharpe", "Sortino", "Beta",
        "Max Drawdown", "VaR 95 %", "VaR 99 %", "CVaR",
        "Risk Contrib.", "Return Contrib.",
    ]
    for col, h in enumerate(headers, 1):
        _hdr_cell(ws, 1, col, h)

    for row_i, m in enumerate(result.asset_metrics.values(), start=2):
        alt = (row_i % 2 == 0)
        cells = [
            (m.ticker,                              _LEFT),
            (_pct(m.weight),                        _RIGHT),
            (_eur(m.latest_price),                  _RIGHT),
            (_pct(m.cagr),                          _RIGHT),
            (_pct(m.annualized_return),             _RIGHT),
            (_pct(m.volatility),                    _RIGHT),
            (_num(m.sharpe),                        _RIGHT),
            (_num(m.sortino),                       _RIGHT),
            (_num(m.beta),                          _RIGHT),
            (_pct(m.max_drawdown),                  _RIGHT),
            (_pct(m.var_95),                        _RIGHT),
            (_pct(m.var_99),                        _RIGHT),
            (_pct(m.cvar),                          _RIGHT),
            (_pct(m.risk_contribution),             _RIGHT),
            (_pct(m.return_contribution),           _RIGHT),
        ]
        for col, (val, align) in enumerate(cells, 1):
            _data_cell(ws, row_i, col, val, align=align, alt=alt)

    _auto_width(ws)


def _build_optimization(ws, result: AnalysisResult) -> None:
    ws.title = "Optimization"
    ws.sheet_view.showGridLines = False

    opt_map = {
        "max_sharpe":      "Maximum Sharpe",
        "min_variance":    "Minimum Variance",
        "black_litterman": "Black-Litterman",
    }
    current = result.optimization.get("current")
    tickers = (
        list(current.weights.keys())
        if current and current.weights
        else list(result.asset_metrics.keys())
    )

    row = 1
    for key, label in opt_map.items():
        opt = result.optimization.get(key)

        # Section heading
        ws.merge_cells(f"A{row}:E{row}")
        heading = ws.cell(row=row, column=1, value=label.upper())
        heading.font  = _LABEL_FONT
        row += 1

        # Weight table
        for col, h in enumerate(["Ticker", "Current Weight", "Optimized Weight"], 1):
            _hdr_cell(ws, row, col, h)
        row += 1
        for i, ticker in enumerate(tickers):
            alt = i % 2 == 1
            cw = current.weights.get(ticker, 0.0) if current else 0.0
            ow = opt.weights.get(ticker, 0.0) if (opt and opt.weights) else 0.0
            _data_cell(ws, row, 1, ticker, alt=alt)
            _data_cell(ws, row, 2, _pct(cw), align=_RIGHT, alt=alt)
            cell = ws.cell(row=row, column=3, value=_pct(ow))
            cell.font      = _BODY_FONT
            cell.border    = _CELL_BORDER
            cell.alignment = _RIGHT
            if alt:
                cell.fill = _ALT_FILL
            row += 1

        row += 1

        # Metric comparison
        for col, h in enumerate(["Metric", "Current", "Optimized"], 1):
            _hdr_cell(ws, row, col, h)
        row += 1
        metric_defs = [
            ("Expected Return", lambda o: _pct(o.expected_return)),
            ("Volatility",      lambda o: _pct(o.volatility)),
            ("Sharpe",          lambda o: _num(o.sharpe)),
            ("Beta",            lambda o: _num(o.beta)),
            ("VaR 95 %",        lambda o: _pct(o.var_95)),
            ("Max Drawdown",    lambda o: _pct(o.max_drawdown)),
        ]
        for i, (mlabel, fn) in enumerate(metric_defs):
            alt = i % 2 == 1
            _data_cell(ws, row, 1, mlabel, alt=alt)
            _data_cell(ws, row, 2, fn(current) if current else "", align=_RIGHT, alt=alt)
            _data_cell(ws, row, 3, fn(opt) if opt else "", align=_RIGHT, alt=alt)
            row += 1

        row += 2  # blank rows between optimizers

    _auto_width(ws)


def _build_scenarios(ws, result: AnalysisResult) -> None:
    ws.title = "Scenarios"
    ws.sheet_view.showGridLines = False

    try:
        weights = {t: m.weight for t, m in result.asset_metrics.items()}
        start_val = float(result.portfolio_value_series.iloc[-1]) \
            if not result.portfolio_value_series.empty else 10_000.0
        scenarios = run_scenario_analysis(
            result.returns, weights, result.benchmark_returns, start_value=start_val,
        )
    except Exception:
        ws.cell(row=1, column=1, value="Scenario data unavailable.")
        return

    headers = [
        "Scenario", "Market Return",
        "Portfolio Return", "New Value (€)",
        "Change vs Baseline", "Stressed Volatility",
    ]
    for col, h in enumerate(headers, 1):
        _hdr_cell(ws, 1, col, h)

    for i, sc in enumerate(scenarios, start=2):
        alt = i % 2 == 0
        change = sc.new_value - start_val
        _data_cell(ws, i, 1, sc.name, alt=alt)
        _data_cell(ws, i, 2, _pct(sc.market_return), align=_RIGHT, alt=alt)
        _data_cell(ws, i, 3, _pct(sc.portfolio_return), align=_RIGHT, alt=alt)
        _data_cell(ws, i, 4, _eur(sc.new_value), align=_RIGHT, alt=alt)
        _data_cell(ws, i, 5, f"{change:+,.0f} €", align=_RIGHT, alt=alt)
        _data_cell(ws, i, 6, _pct(sc.stressed_vol), align=_RIGHT, alt=alt)

    _auto_width(ws)


# ── Public API ─────────────────────────────────────────────────────────────────

def export_portfolio_excel(result: AnalysisResult, path: Path) -> None:
    """Write a professionally styled multi-sheet Excel workbook.

    Sheets: Summary | Assets | Optimization | Scenarios

    Args:
        result: Full AnalysisResult from the analysis engine.
        path:   Destination .xlsx file path.
    """
    wb = openpyxl.Workbook()
    # Remove the default empty sheet
    wb.remove(wb.active)

    _build_summary(wb.create_sheet(), result)
    _build_assets(wb.create_sheet(), result)
    _build_optimization(wb.create_sheet(), result)
    _build_scenarios(wb.create_sheet(), result)

    wb.save(str(path))
